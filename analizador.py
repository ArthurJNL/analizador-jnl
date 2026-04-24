import streamlit as st
import pandas as pd
from datetime import datetime
import math
import uuid

# --- MOTORES EXTERNOS ---
try:
    from docx import Document
except ImportError:
    Document = None

try:
    from st_keyup import st_keyup
except ImportError:
    st_keyup = None

try:
    from streamlit_echarts import st_echarts
except ImportError:
    st_echarts = None

try:
    from fpdf import FPDF
except ImportError:
    FPDF = None

# 1. CONFIGURAÇÕES DA PÁGINA
st.set_page_config(page_title="Analisador JNL", page_icon="🛡️", layout="wide")

# --- DESIGN ESTÁVEL E MINIMALISTA ---
st.markdown("""
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap');
    
    html, body, [class*="css"] { 
        font-family: 'Inter', sans-serif; 
    }
    
    .stApp { background-color: #F8F9FA; }
    
    .stTextInput input {
        border-radius: 8px !important; 
        border: 1px solid #D0D5DD !important; 
    }
    .stTextInput input:focus {
        border-color: #111111 !important; 
        box-shadow: 0 0 0 1px #111111 !important;
    }
    
    .stDeployButton {display:none;}
    #MainMenu {visibility: hidden;}
    footer {visibility: hidden;}
    </style>
    """, unsafe_allow_html=True)

st.title("🛡️ ANALISADOR JNL")
st.write("Análise inteligente e controle operacional.")
st.markdown("---")

# --- MOTOR DE PESQUISA INTELIGENTE (KEYUP) ---
def campo_pesquisa(label, placeholder, key):
    if st_keyup is not None:
        return st_keyup(label, placeholder=placeholder, key=key)
    else:
        return st.text_input(label, placeholder=placeholder, key=key)

# --- BARRA LATERAL ---
with st.sidebar:
    st.header("⚙️ Operação")
    st.markdown("---")
    arquivos_enviados = st.file_uploader("Arraste seus documentos e planilhas", type=["xlsx", "xls", "xlsm", "docx", "txt", "csv"], accept_multiple_files=True)

def formatar_moeda(valor):
    try:
        if pd.isna(valor): return "R$ 0,00"
        return f"R$ {float(valor):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    except: return "R$ 0,00"

def formatar_orcamento(valor):
    if pd.isna(valor) or str(valor).strip() == "" or str(valor).strip().upper() in ["S/N", "NAN", "NONE"]:
        return "S/N"
    try: return str(int(float(valor)))
    except: return str(valor).strip()

# ==========================================
# MOTORES DE RELATÓRIO PDF
# ==========================================
def limpar_texto(t):
    import unicodedata
    return unicodedata.normalize('NFKD', str(t)).encode('ASCII', 'ignore').decode('utf-8')

if FPDF is not None:
    class PDFReport(FPDF):
        def header(self):
            self.set_font('Arial', 'B', 10)
            self.cell(0, 10, 'ANALISADOR JNL - Relatorio Oficial', 0, 1, 'C')
            self.ln(2)
        def footer(self):
            self.set_y(-15)
            self.set_font('Arial', 'I', 8)
            self.cell(0, 10, f'Pagina {self.page_no()}', 0, 0, 'C')

    def obter_larguras_dinamicas(colunas):
        page_width = 190
        widths = []
        for col in colunas:
            c = col.upper()
            if 'DESCRI' in c or 'RAZAO' in c or 'NOME' in c or 'ITEM' in c: widths.append(page_width * 0.35)
            elif 'OBS' in c: widths.append(page_width * 0.25)
            elif 'DATA' in c or 'SITUA' in c or 'MARCA' in c or 'LOCAL' in c or 'PRAT' in c: widths.append(page_width * 0.15)
            elif 'QTD' in c or 'QUANTIDADE' in c or 'MIN' in c or 'MÍN' in c: widths.append(page_width * 0.1)
            else: widths.append(page_width * 0.15)
        
        total = sum(widths)
        return [w * (page_width / total) for w in widths]

    def gerar_pdf_tabela(df, titulo):
        pdf = PDFReport()
        pdf.add_page()
        pdf.set_font("Arial", 'B', 12)
        pdf.cell(0, 10, limpar_texto(titulo), 0, 1, 'C')
        pdf.ln(5)
        
        colunas = list(df.columns)
        widths = obter_larguras_dinamicas(colunas)
            
        pdf.set_fill_color(17, 17, 17)
        pdf.set_text_color(255, 255, 255)
        pdf.set_font("Arial", 'B', 9)
        
        for i, col in enumerate(colunas):
            pdf.cell(widths[i], 8, limpar_texto(col), border=1, fill=True, align='C')
        pdf.ln()
        
        line_height = 5
        pdf.set_font("Arial", '', 8)
        pdf.set_fill_color(255, 255, 255)
        pdf.set_text_color(26, 28, 30)
        
        for _, row in df.iterrows():
            max_linhas = 1
            for i, item in enumerate(row):
                texto = limpar_texto(item)
                w_util = widths[i] - 2
                w_texto = pdf.get_string_width(texto)
                linhas = math.ceil(w_texto / w_util) if w_util > 0 else 1
                if linhas > max_linhas:
                    max_linhas = linhas
                    
            h_linha = (max_linhas * line_height) + 2
            
            if pdf.get_y() + h_linha > 275:
                pdf.add_page()
                pdf.set_fill_color(17, 17, 17)
                pdf.set_text_color(255, 255, 255)
                pdf.set_font("Arial", 'B', 9)
                for i, col in enumerate(colunas):
                    pdf.cell(widths[i], 8, limpar_texto(col), border=1, fill=True, align='C')
                pdf.ln()
                pdf.set_font("Arial", '', 8)
                pdf.set_fill_color(255, 255, 255)
                pdf.set_text_color(26, 28, 30)
                    
            start_x = pdf.get_x()
            start_y = pdf.get_y()
            
            for i, item in enumerate(row):
                texto = limpar_texto(item)
                w = widths[i]
                x = start_x + sum(widths[:i])
                y = start_y
                
                pdf.rect(x, y, w, h_linha, 'D')
                pdf.set_xy(x, y + 1)
                
                align = 'L' if i == 0 else 'C'
                pdf.multi_cell(w, line_height, texto, border=0, align=align)
                
            pdf.set_xy(start_x, start_y + h_linha)
            
        res = pdf.output(dest='S')
        if isinstance(res, str): return res.encode('latin-1')
        return bytes(res)

    def gerar_pdf_ranking(df, titulo, tipo="estoque"):
        pdf = PDFReport()
        pdf.add_page()
        pdf.set_font("Arial", 'B', 12)
        pdf.cell(0, 10, limpar_texto(titulo), 0, 1, 'C')
        pdf.ln(5)
        
        pdf.set_fill_color(17, 17, 17)
        pdf.set_text_color(255, 255, 255)
        pdf.set_font("Arial", 'B', 9)
        widths = [20, 120, 50]
        
        if tipo == "financeiro": colunas = ["POS.", "RAZAO SOCIAL / DESCRICAO", "VALOR TOTAL"]
        else: colunas = ["POS.", "ITEM / DESCRICAO", "QUANTIDADE EM ESTOQUE"]
            
        for i, col in enumerate(colunas):
            pdf.cell(widths[i], 8, col, border=1, fill=True, align='C')
        pdf.ln()
        
        pdf.set_text_color(26, 28, 30)
        pdf.set_font("Arial", '', 8)
        line_height = 5
        
        col_nome = df.columns[0]
        col_valor = df.columns[1]
        df_ord = df.sort_values(by=col_valor, ascending=False).reset_index(drop=True)
        
        for i, row in df_ord.iterrows():
            pos = f"{i + 1}."
            nome = limpar_texto(row[col_nome]) 
            if tipo == "financeiro": valor = f"R$ {row[col_valor]:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
            else: valor = f"{int(row[col_valor])} un."
                
            linha_dados = [pos, nome, valor]
            max_linhas = 1
            for j, item in enumerate(linha_dados):
                w_util = widths[j] - 2
                w_texto = pdf.get_string_width(item)
                linhas = math.ceil(w_texto / w_util) if w_util > 0 else 1
                if linhas > max_linhas: max_linhas = linhas
                    
            h_linha = (max_linhas * line_height) + 2
            
            if pdf.get_y() + h_linha > 275:
                pdf.add_page()
                pdf.set_fill_color(17, 17, 17)
                pdf.set_text_color(255, 255, 255)
                pdf.set_font("Arial", 'B', 9)
                for j, col in enumerate(colunas): pdf.cell(widths[j], 8, col, border=1, fill=True, align='C')
                pdf.ln()
                pdf.set_text_color(26, 28, 30)
                pdf.set_font("Arial", '', 8)
                
            start_x = pdf.get_x()
            start_y = pdf.get_y()
            
            for j, item in enumerate(linha_dados):
                w = widths[j]
                x = start_x + sum(widths[:j])
                y = start_y
                pdf.rect(x, y, w, h_linha, 'D')
                pdf.set_xy(x, y + 1)
                align = 'C' if j == 0 else ('L' if j == 1 else 'R')
                pdf.multi_cell(w, line_height, item, border=0, align=align)
                
            pdf.set_xy(start_x, start_y + h_linha)
            
        res = pdf.output(dest='S')
        if isinstance(res, str): return res.encode('latin-1')
        return bytes(res)

# ==========================================
# AGENDA ICS
# ==========================================
def criar_lembrete_item(data_venc, cliente, valor, orc):
    if pd.isnull(data_venc): return None
    dtstart = data_venc.strftime("%Y%m%d") + "T100000"
    dtend = data_venc.strftime("%Y%m%d") + "T103000"
    cliente_limpo = str(cliente).replace("\n", " ")
    uid = f"{uuid.uuid4()}@jnl.com"
    return f"BEGIN:VCALENDAR\nVERSION:2.0\nPRODID:-//JNL//Lembrete Financeiro//PT\nBEGIN:VEVENT\nUID:{uid}\nDTSTART:{dtstart}\nDTEND:{dtend}\nSUMMARY:⚠️ COBRAR: {cliente_limpo} ({formatar_moeda(valor)})\nDESCRIPTION:Lembrete JNL\\nOrçamento: {formatar_orcamento(orc)}\\nValor: {formatar_moeda(valor)}\nBEGIN:VALARM\nTRIGGER:-P1D\nACTION:DISPLAY\nDESCRIPTION:Vence Amanhã\nEND:VALARM\nEND:VEVENT\nEND:VCALENDAR"

def criar_lembrete_estoque(item, qtd, minimo):
    hoje = datetime.now()
    dtstart = hoje.strftime("%Y%m%d") + "T090000"
    dtend = hoje.strftime("%Y%m%d") + "T093000"
    item_limpo = str(item).replace("\n", " ")
    uid = f"{uuid.uuid4()}@jnl.com"
    return f"BEGIN:VCALENDAR\nVERSION:2.0\nPRODID:-//JNL//Lembrete Estoque//PT\nBEGIN:VEVENT\nUID:{uid}\nDTSTART:{dtstart}\nDTEND:{dtend}\nSUMMARY:📦 REPOR ESTOQUE: {item_limpo}\nDESCRIPTION:Alerta JNL\\nO item '{item_limpo}' atingiu nível crítico.\\nSaldo Atual: {qtd} unidades\\nMínimo Exigido: {minimo} unidades\\nNecessário solicitar reposição imediatamente.\nBEGIN:VALARM\nTRIGGER:-PT5M\nACTION:DISPLAY\nDESCRIPTION:Reposição de Estoque\nEND:VALARM\nEND:VEVENT\nEND:VCALENDAR"

# ==========================================
# PROCESSAMENTO DE ARQUIVOS
# ==========================================
if arquivos_enviados:
    for arquivo in arquivos_enviados:
        extensao = arquivo.name.split('.')[-1].lower()
        
        with st.expander(f"📄 DOCUMENTO ATIVO: {arquivo.name.upper()}", expanded=True):
            
            if extensao in ['xlsx', 'xls', 'xlsm', 'csv']:
                try:
                    if extensao == 'csv': df = pd.read_csv(arquivo, sep=None, engine='python')
                    else: df = pd.read_excel(arquivo)
                    
                    if any("Unnamed" in str(c) for c in df.columns):
                        for idx, row in df.head(15).iterrows():
                            linha_texto = " ".join([str(x).lower() for x in row.values])
                            if any(k in linha_texto for k in ['valor', 'r$', 'data', 'venc', 'cliente', 'item', 'qtd', 'estoque', 'marca', 'prateleira']):
                                nomes_limpos = []
                                for i, c in enumerate(row.values):
                                    nome = str(c).strip() if pd.notna(c) and str(c).strip() != "" else f"vazio_{i}"
                                    if nome in nomes_limpos: nome = f"{nome}_{i}"
                                    nomes_limpos.append(nome)
                                df.columns = nomes_limpos
                                df = df.iloc[idx+1:].reset_index(drop=True)
                                break
                    
                    df = df.dropna(how='all', axis=1).dropna(how='all', axis=0)
                    cols_limpas = {str(c).lower().strip(): c for c in df.columns}

                    is_financeiro = any(k in str(cols_limpas.keys()) for k in ['vencimento', 'data']) and any(k in str(cols_limpas.keys()) for k in ['valor', 'r$'])

                    # --- FLUXO 1: FINANCEIRO ---
                    if is_financeiro:
                        st.info("🎯 **Módulo Financeiro Ativado**")
                        col_data = next((v for k, v in cols_limpas.items() if 'vencimento' in k or 'data' in k), None)
                        col_valor = next((v for k, v in cols_limpas.items() if 'valor' in k or 'r$' in k), None)
                        col_cliente = next((v for k, v in cols_limpas.items() if 'cliente' in k or 'nome' in k or 'empresa' in k), "S/N")
                        col_orc = next((v for k, v in cols_limpas.items() if 'orc' in k or 'orç' in k or 'pedido' in k or 'doc' in k), None)
                        col_status = next((v for k, v in cols_limpas.items() if 'obs' in k or 'status' in k or 'situa' in k), None)
                        
                        df[col_data] = pd.to_datetime(df[col_data], errors='coerce')
                        df[col_valor] = pd.to_numeric(df[col_valor], errors='coerce').fillna(0)
                        
                        mask_pago = df[col_status].astype(str).str.lower().str.contains('pago') if col_status else df[df.columns[-1]].astype(str).str.lower().str.contains('pago')
                        df_pendente = df[~mask_pago].dropna(subset=[col_data])
                        
                        hoje = pd.to_datetime('today').normalize()
                        df_a_vencer = df_pendente[df_pendente[col_data] >= hoje].sort_values(by=col_data)

                        tab_venc, tab_dados = st.tabs(["📅 Resumo Financeiro", "📋 Base de Dados"])
                        with tab_venc:
                            st.markdown(f"### Valor total em aberto: {formatar_moeda(df_pendente[col_valor].sum())}")
                            for _, linha in df_a_vencer.head(10).iterrows():
                                c1, c2 = st.columns([0.85, 0.15])
                                c1.write(f"📌 **{linha.get(col_cliente, 'S/N')}** | {formatar_moeda(linha[col_valor])} | Venc: {linha[col_data].strftime('%d/%m/%Y')}")
                                c2.download_button("🔔 Lembrete", data=criar_lembrete_item(linha[col_data], linha.get(col_cliente, 'S/N'), linha[col_valor], linha.get(col_orc)), file_name="Lembrete.ics", key=str(uuid.uuid4()))
                        
                        with tab_dados:
                            busca_fin = campo_pesquisa("🔍 Pesquisar no Financeiro", "Digite cliente, nota ou valor para filtrar instantaneamente...", key=f"bf_{arquivo.name}")
                            df_view = df.copy()
                            if busca_fin:
                                mask_fin = df_view.astype(str).apply(lambda x: x.str.contains(busca_fin, case=False, na=False)).any(axis=1)
                                df_view = df_view[mask_fin]
                            st.dataframe(df_view, use_container_width=True)

                    # --- FLUXO 2: PATRIMÔNIO / INVENTÁRIO ---
                    else:
                        st.info("🎯 **Módulo de Logística e Inventário**")
                        
                        col_qtd = next((v for k, v in cols_limpas.items() if any(x in k for x in ['qtd', 'quantidade', 'saldo', 'estoque'])), None)
                        col_desc = next((v for k, v in cols_limpas.items() if any(x in k for x in ['descri', 'item', 'produto', 'nome'])), None)
                        col_marca = next((v for k, v in cols_limpas.items() if 'marca' in k or 'fabricante' in k), None)
                        col_prat = next((v for k, v in cols_limpas.items() if 'prateleira' in k or 'local' in k), None)
                        col_obs = next((v for k, v in cols_limpas.items() if 'obs' in k or 'coment' in k), None)
                        col_minimo = next((v for k, v in cols_limpas.items() if 'mínimo' in k or 'minimo' in k), None)

                        if col_qtd and col_desc:
                            df[col_qtd] = pd.to_numeric(df[col_qtd], errors='coerce').fillna(0)
                            col_m1, col_m2, col_m3 = st.columns(3)
                            
                            if col_minimo:
                                df[col_minimo] = pd.to_numeric(df[col_minimo], errors='coerce').fillna(0)
                                df_critico = df[df[col_qtd] <= df[col_minimo]].sort_values(by=col_qtd)
                                
                                col_m1.metric("Itens Analisados", len(df))
                                col_m2.metric("Total geral de itens", f"{int(df[col_qtd].sum())} un.")
                                col_m3.metric("🚨 Itens para repor", len(df_critico))

                                if not df_critico.empty:
                                    st.error(f"**Atenção:** {len(df_critico)} SKUs atingiram a cota mínima de reposição.")
                                    for _, linha in df_critico.iterrows():
                                        c1, c2 = st.columns([0.85, 0.15])
                                        item_nome = linha[col_desc]
                                        saldo = int(linha[col_qtd])
                                        minimo_item = int(linha[col_minimo])
                                        c1.write(f"📦 **{item_nome}** | Saldo: **{saldo}** (Mín: **{minimo_item}**) | Marca: {linha.get(col_marca, 'S/M')}")
                                        c2.download_button(label="🔔 Repor", data=criar_lembrete_estoque(item_nome, saldo, minimo_item), file_name=f"Repor_{str(item_nome)[:10]}.ics", key=str(uuid.uuid4()))
                            else:
                                col_m1.metric("Itens Analisados", len(df))
                                col_m2.metric("Total geral de itens", f"{int(df[col_qtd].sum())} un.")
                                st.warning("⚠️ **Atenção:** Coluna de 'ESTOQUE MÍNIMO' não detectada no arquivo.")

                        st.write("---")
                        busca_est = campo_pesquisa("🔍 Filtro", "Pesquise por descrição, marca ou prateleira...", key=f"be_{arquivo.name}")
                        
                        colunas_desejadas = [c for c in [col_qtd, col_minimo, col_desc, col_marca, col_prat, col_obs] if c]
                        
                        if colunas_desejadas:
                            df_filtrado = df[colunas_desejadas].copy()
                            if busca_est:
                                mask_est = df_filtrado.astype(str).apply(lambda x: x.str.contains(busca_est, case=False, na=False)).any(axis=1)
                                df_filtrado = df_filtrado[mask_est]
                            
                            # 💡 ABAS INVERTIDAS: Tabela primeiro, Gráfico depois!
                            aba_tab, aba_visu = st.tabs(["📋 Base de Dados Operacional", "📊 Distribuição de Estoque"])
                            
                            with aba_tab:
                                titulo_tab_est = st.text_input("📝 Título do Relatório PDF (Tabela):", value=f"Controle de Estoque - {datetime.now().strftime('%d/%m/%Y')}", key=f"tt_{arquivo.name}")
                                nomes_exibicao = {col_qtd: "QUANTIDADE", col_minimo: "ESTOQUE MÍNIMO", col_desc: "DESCRIÇÃO", col_marca: "MARCA", col_prat: "PRATELEIRA", col_obs: "OBSERVAÇÕES"}
                                df_tabela = df_filtrado.rename(columns={k: v for k, v in nomes_exibicao.items() if k in df_filtrado.columns})
                                
                                col_t1, col_t2 = st.columns([0.8, 0.2])
                                with col_t1: st.write("💡 *Baixe em PNG ou PDF.*")
                                with col_t2:
                                    if FPDF is not None and not df_tabela.empty:
                                        st.download_button(label="📄 Baixar PDF", data=gerar_pdf_tabela(df_tabela, titulo_tab_est), file_name=f"Tabela_Estoque.pdf", mime="application/pdf", key=f"btn_tb_{arquivo.name}", use_container_width=True)
                                    elif FPDF is None:
                                        st.error("⚠️ Falta a biblioteca 'fpdf' no GitHub!")

                                st.dataframe(df_tabela, use_container_width=True)

                            with aba_visu:
                                mostrar_grafico = st.toggle("Exibir Gráfico de Volumetria", value=False, key=f"tgl_graf_{arquivo.name}")
                                
                                if mostrar_grafico:
                                    titulo_graf_est = st.text_input("📝 Título do Relatório PDF (Gráfico):", value=f"CONTROLE DE VOLUME - {datetime.now().strftime('%d/%m/%Y')}", key=f"tg_{arquivo.name}")
                                    dados_grafico = df_filtrado.groupby(col_desc)[col_qtd].sum().reset_index().sort_values(by=col_qtd, ascending=True)
                                    
                                    col_g1, col_g2 = st.columns([0.8, 0.2])
                                    with col_g1: st.write("💡 *Baixe em PNG ou PDF.*")
                                    with col_g2:
                                        if FPDF is not None and not dados_grafico.empty:
                                            st.download_button(label="📄 Baixar PDF", data=gerar_pdf_ranking(dados_grafico, titulo_graf_est, tipo="estoque"), file_name=f"Ranking_Estoque.pdf", mime="application/pdf", key=f"btn_rk_{arquivo.name}", use_container_width=True)
                                        elif FPDF is None:
                                            st.error("⚠️ Falta a biblioteca 'fpdf' no GitHub!")

                                    if not df_filtrado.empty and st_echarts is not None:
                                        dados_barras_formatados = [{"value": int(row[col_qtd]), "label": {"show": True, "position": "right", "formatter": "{c} un.", "color": "#111111", "fontWeight": "bold"}} for _, row in dados_grafico.iterrows()]
                                        altura_dinamica = max(500, len(dados_grafico) * 45)
                                        
                                        bar_options = {
                                            "backgroundColor": "transparent",
                                            "title": {"text": "Volumetria por SKU", "left": "center", "textStyle": {"color": "#111111", "fontSize": 16, "fontFamily": "Inter"}},
                                            "toolbox": {"feature": {"saveAsImage": {"show": True, "title": "Baixar PNG", "pixelRatio": 2}}},
                                            "tooltip": {"trigger": "axis", "axisPointer": {"type": "shadow"}},
                                            "grid": {"top": 60, "left": "1%", "right": "10%", "bottom": "1%", "containLabel": True},
                                            "xAxis": {"type": "value", "splitLine": {"lineStyle": {"type": "dashed", "color": "#E2E8F0"}}},
                                            "yAxis": {"type": "category", "data": dados_grafico[col_desc].tolist(), "axisLabel": {"interval": 0, "width": 250, "overflow": "break", "lineHeight": 14, "color": "#111111"}},
                                            "series": [{"type": "bar", "data": dados_barras_formatados, "itemStyle": {"color": "#111111", "borderRadius": [0, 6, 6, 0]}}]
                                        }
                                        st_echarts(options=bar_options, height=f"{altura_dinamica}px")
                                else:
                                    st.info("📊 O gráfico está desativado para otimizar o desempenho. Clique no interruptor acima para visualizá-lo.")

                except Exception as e:
                    st.error(f"Erro no processamento do arquivo: {e}")

            elif extensao in ['docx', 'txt']:
                st.info("🎯 **Módulo de Documentação Ativado**")
                if Document is None and extensao == 'docx':
                    st.error("Biblioteca 'python-docx' ausente no servidor.")
                else:
                    conteudo = [p.text for p in Document(arquivo).paragraphs] if extensao == 'docx' else arquivo.read().decode("utf-8").splitlines()
                    busca = campo_pesquisa(f"🔍 Busca Dinâmica", "Pesquise palavras-chave no documento...", key=f"bd_{arquivo.name}")
                    if busca:
                        resultados = [l for l in conteudo if busca.lower() in l.lower()]
                        if resultados:
                            for r in resultados: st.write(f"🔹 {r}")
                        else:
                            st.warning("Nenhum trecho correspondente encontrado.")

else:
    st.info("Aguardando inserção de dados para iniciar os motores de análise.")